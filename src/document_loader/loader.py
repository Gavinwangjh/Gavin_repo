import re
import json
import hashlib
from pathlib import Path
from typing import List, Dict, Any, Optional, Callable
from dataclasses import dataclass

import PyPDF2
from docx import Document as DocxDocument
from loguru import logger

from .article_chunker import ArticleChunker


@dataclass
class Document:
    """文档数据结构"""
    content: str
    metadata: Dict[str, Any]
    doc_id: str

    def __post_init__(self):
        if not self.doc_id:
            self.doc_id = self._generate_doc_id()

    def _generate_doc_id(self) -> str:
        content_hash = hashlib.md5(self.content.encode("utf-8")).hexdigest()
        return f"doc_{content_hash[:16]}"


@dataclass
class DocumentChunk:
    """文档分块数据结构"""
    content: str
    metadata: Dict[str, Any]
    chunk_id: str
    doc_id: str
    chunk_index: int

    def __post_init__(self):
        if not self.chunk_id:
            self.chunk_id = f"{self.doc_id}_chunk_{self.chunk_index}"


class RecursiveTextSplitter:
    """
    递归文本切割（语义优先）
    分割顺序：段落 -> 句子 -> 标点 -> 硬切
    目标：尽量不把语义切碎，同时保证 chunk_size。
    """

    def __init__(
        self,
        chunk_size: int = 512,
        chunk_overlap: int = 50,
        min_chunk_size: int = 120,  # ✅ 添加这个参数
        separators: Optional[List[str]] = None,
    ):
        self.chunk_size = chunk_size
        self.chunk_overlap = max(0, min(chunk_overlap, chunk_size // 2))
        self.min_chunk_size = min_chunk_size  # ✅ 添加这个赋值
        self.separators = separators or [
            "\n\n",          # 段落
            "\n",            # 换行
            "。", "！", "？", "；",  # 句末
            "，", "、",       # 逗号/顿号
        ]
    
    # ... 其余代码不变

    def split(self, text: str) -> List[str]:
        text = (text or "").strip()
        if not text:
            return []

        # 递归拆
        pieces = self._recursive_split(text, self.separators)

        # 组装到 chunk_size，并做 overlap
        chunks = self._merge_pieces(pieces)

        # 合并过短块（避免碎片）
        chunks = self._merge_too_small(chunks)

        return chunks

    def _recursive_split(self, text: str, seps: List[str]) -> List[str]:
        if len(text) <= self.chunk_size:
            return [text]

        if not seps:
            # 没有分隔符了，硬切
            return [text[i:i + self.chunk_size] for i in range(0, len(text), self.chunk_size)]

        sep = seps[0]

        if sep in ["。", "！", "？", "；", "，", "、"]:
            # 保留分隔符（更像句子/短语）
            parts = re.split(f"({re.escape(sep)})", text)
            merged = []
            buf = ""
            for p in parts:
                if not p:
                    continue
                buf += p
                # 句号类：遇到标点就提交一次（更自然）
                if p == sep:
                    merged.append(buf)
                    buf = ""
            if buf.strip():
                merged.append(buf)
        else:
            merged = text.split(sep)

        # 如果分割没有起效，继续用下一个 sep
        if len(merged) <= 1:
            return self._recursive_split(text, seps[1:])

        out: List[str] = []
        for part in merged:
            part = part.strip()
            if not part:
                continue
            if len(part) <= self.chunk_size:
                out.append(part)
            else:
                out.extend(self._recursive_split(part, seps[1:]))

        return out

    def _merge_pieces(self, pieces: List[str]) -> List[str]:
        chunks: List[str] = []
        buf = ""

        for p in pieces:
            p = p.strip()
            if not p:
                continue

            if not buf:
                buf = p
                continue

            # +1 预留一个分隔空格
            if len(buf) + 1 + len(p) <= self.chunk_size:
                buf = f"{buf} {p}"
            else:
                chunks.append(buf.strip())
                buf = p

        if buf.strip():
            chunks.append(buf.strip())

        # overlap（按字符简单实现）
        if self.chunk_overlap <= 0 or len(chunks) <= 1:
            return chunks

        overlapped: List[str] = []
        for i, c in enumerate(chunks):
            if i == 0:
                overlapped.append(c)
                continue
            prev = overlapped[-1]
            tail = prev[-self.chunk_overlap:] if len(prev) > self.chunk_overlap else prev
            overlapped.append((tail + c).strip())

        return overlapped

    def _merge_too_small(self, chunks: List[str]) -> List[str]:
        if not chunks:
            return []
        merged: List[str] = []
        i = 0
        while i < len(chunks):
            cur = chunks[i]
            if len(cur) >= self.min_chunk_size:
                merged.append(cur)
                i += 1
                continue

            # 太短：尽量跟后面合并
            if i + 1 < len(chunks):
                nxt = chunks[i + 1]
                combined = (cur + "\n" + nxt).strip()
                merged.append(combined)
                i += 2
            else:
                # 最后一个短块：塞到上一个
                if merged:
                    merged[-1] = (merged[-1] + "\n" + cur).strip()
                else:
                    merged.append(cur)
                i += 1
        return merged


class DocumentLoader:
    """文档加载器，支持多种格式"""

    SUPPORTED_FORMATS = {".pdf", ".txt", ".md", ".docx"}

    def __init__(
        self,
        chunk_size: int = 512,
        chunk_overlap: int = 50,
        use_article_chunking: bool = True,
    ):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.use_article_chunking = use_article_chunking

        if self.use_article_chunking:
            self.article_chunker = ArticleChunker(
                max_chunk_size=chunk_size * 3,
                min_chunk_size=max(80, chunk_size // 4),
            )

        self.recursive_splitter = RecursiveTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            min_chunk_size=max(120, chunk_size // 3),
        )

    def load_document(self, file_path: str) -> Document:
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        if path.suffix.lower() not in self.SUPPORTED_FORMATS:
            raise ValueError(f"不支持的文件格式: {path.suffix}")

        logger.info(f"加载文档: {file_path}")

        if path.suffix.lower() == ".pdf":
            content = self._load_pdf(path)
        elif path.suffix.lower() == ".docx":
            content = self._load_docx(path)
        else:
            content = self._load_text(path)

        metadata = {
            "filename": path.name,
            "file_path": str(path.absolute()),
            "file_size": path.stat().st_size,
            "file_type": path.suffix.lower(),
            "created_time": path.stat().st_ctime,
            "modified_time": path.stat().st_mtime,
        }

        doc = Document(content=content, metadata=metadata, doc_id="")
        logger.info(f"文档加载成功: {path.name}, 内容长度: {len(content)}")
        return doc

    def _load_pdf(self, path: Path) -> str:
        content = []
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for idx, page in enumerate(reader.pages):
                try:
                    txt = page.extract_text() or ""
                    if txt.strip():
                        content.append(txt)
                except Exception as e:
                    logger.warning(f"PDF第{idx+1}页解析失败: {e}")
        if not content:
            raise ValueError("PDF文件没有可提取的文本内容")
        return self._clean_text("\n".join(content))

    def _load_docx(self, path: Path) -> str:
        doc = DocxDocument(path)
        content = []
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                content.append(t)
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    ct = (cell.text or "").strip()
                    if ct:
                        row_text.append(ct)
                if row_text:
                    content.append(" | ".join(row_text))
        if not content:
            raise ValueError("DOCX文件没有可提取的文本内容")
        return self._clean_text("\n".join(content))

    def _load_text(self, path: Path) -> str:
        encodings = ["utf-8", "gbk", "gb2312", "ascii"]
        last_err = None
        for enc in encodings:
            try:
                with open(path, "r", encoding=enc) as f:
                    content = f.read()
                if content.strip():
                    return self._clean_text(content)
                raise ValueError("文件内容为空")
            except Exception as e:
                last_err = e
        raise ValueError(f"文本文件读取失败: {last_err}")

    def _clean_text(self, text: str) -> str:
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\r", "\n", text)

        lines = text.split("\n")
        cleaned_lines = [re.sub(r"\s+", " ", ln).strip() for ln in lines if ln.strip()]
        structured_patterns = [
            r"^第[一二三四五六七八九十百千万\d]+[章节条款项]",
            r"^[一二三四五六七八九十]+、",
            r"^（[一二三四五六七八九十\d]+）",
            r"^\d+[、.]",
        ]
        has_structured_lines = any(
            any(re.match(pattern, line) for pattern in structured_patterns)
            for line in cleaned_lines
        )
        text = ("\n" if has_structured_lines else " ").join(cleaned_lines)

        text = re.sub(r"[^\w\s\u4e00-\u9fff，。！？；：、\"'()[\]《》\-.,!?:;\'\"\n]", "", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def get_file_info(self, file_path: str) -> Dict[str, Any]:
        """Return basic file metadata used by callers and tests."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        return {
            "filename": path.name,
            "file_path": str(path.absolute()),
            "file_size": path.stat().st_size,
            "file_type": path.suffix.lower(),
            "supported": path.suffix.lower() in self.SUPPORTED_FORMATS,
            "is_supported": path.suffix.lower() in self.SUPPORTED_FORMATS,
            "created_time": path.stat().st_ctime,
            "modified_time": path.stat().st_mtime,
        }

    def chunk_document(self, document: Document) -> List[DocumentChunk]:
        content = document.content or ""
        if not content.strip():
            return []

        # 1) 条文优先（如果判定为条文文档）
        if self.use_article_chunking and hasattr(self, "article_chunker"):
            if self.article_chunker.is_article_document(content):
                logger.info(f"检测到条文文档，使用条文分块: {document.doc_id}")
                return self._chunk_by_articles(document)

        # 2) 否则走递归切割（替代原先 size_based）
        logger.info(f"非条文文档，使用递归分块: {document.doc_id}")
        return self._chunk_by_recursive(document)

    def _chunk_by_articles(self, document: Document) -> List[DocumentChunk]:
        article_chunks = self.article_chunker.chunk_by_articles(document.content)
        chunks: List[DocumentChunk] = []
        for i, chunk_data in enumerate(article_chunks):
            md = dict(document.metadata)
            md.update({
                "chunk_index": i,
                "chunk_type": chunk_data.get("chunk_type", "article_based"),
                "strategy": "article",
                "article_number": chunk_data.get("article_number"),
                "article_title": chunk_data.get("article_title"),
                "article_level": chunk_data.get("article_level"),
                "parent_article": chunk_data.get("parent_article"),
                "section_path": chunk_data.get("section_path"),
                "chunk_type_name": chunk_data.get("chunk_type_name"),
                "chunk_length": len(chunk_data["content"]),
                "chunk_char_start": chunk_data.get("metadata", {}).get("start_pos"),
                "chunk_char_end": chunk_data.get("metadata", {}).get("end_pos"),
            })
            md.update(chunk_data.get("metadata", {}))
            md["chunk_hash"] = hashlib.md5(chunk_data["content"].encode("utf-8")).hexdigest()

            chunks.append(DocumentChunk(
                content=chunk_data["content"],
                metadata=md,
                chunk_id="",
                doc_id=document.doc_id,
                chunk_index=i,
            ))
        return chunks

    def _chunk_by_recursive(self, document: Document) -> List[DocumentChunk]:
        parts = self.recursive_splitter.split(document.content)
        chunks: List[DocumentChunk] = []
        cursor = 0  # 粗略字符定位（用于调试/报表）
        for i, text in enumerate(parts):
            start = document.content.find(text, cursor)
            if start < 0:
                start = cursor
            end = start + len(text)
            cursor = end

            md = dict(document.metadata)
            md.update({
                "chunk_index": i,
                "chunk_type": "recursive",
                "strategy": "recursive",
                "chunk_length": len(text),
                "chunk_char_start": start,
                "chunk_char_end": end,
            })
            md["chunk_hash"] = hashlib.md5(text.encode("utf-8")).hexdigest()

            chunks.append(DocumentChunk(
                content=text,
                metadata=md,
                chunk_id="",
                doc_id=document.doc_id,
                chunk_index=i,
            ))
        logger.info(f"递归分块完成: {document.doc_id}, 共 {len(chunks)} 块")
        return chunks
    
    def load_and_chunk_document(self, file_path: str) -> List[DocumentChunk]:
        """加载并分块文档（CLI upload 依赖）"""
        document = self.load_document(file_path)
        return self.chunk_document(document)

    
